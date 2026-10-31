"""Local document extraction shared by conversion and comparison commands."""

from __future__ import annotations

import json
from email import policy
from email.parser import BytesParser
from io import BytesIO
from pathlib import Path
from typing import cast
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation
from pptx.shapes.autoshape import Shape
from pptx.shapes.graphfrm import GraphicFrame
from pypdf import PdfReader
from pypdf.errors import PdfReadError

PLAIN_TEXT_EXTENSIONS = {
    ".txt", ".md", ".rst", ".csv", ".tsv", ".json", ".xml", ".html", ".htm",
    ".log", ".rss", ".atom",
}
OFFICE_EXTENSIONS = {".docx", ".xlsx", ".pptx"}
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tiff", ".webp"}
SUPPORTED_EXTENSIONS = PLAIN_TEXT_EXTENSIONS | OFFICE_EXTENSIONS | IMAGE_EXTENSIONS | {
    ".pdf", ".ipynb", ".eml", ".zip",
}
CONVERSION_ERRORS = (OSError, ValueError, TypeError, KeyError, BadZipFile, PdfReadError, SyntaxError)
MAX_ARCHIVE_BYTES = 64 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 1000


def extract_docx(data: bytes) -> str:
    document = Document(BytesIO(data))
    parts: list[str] = []
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            parts.append(block.text)
        else:
            for row in block.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def extract_xlsx(data: bytes) -> str:
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=False)
    try:
        parts: list[str] = []
        for sheet in workbook.worksheets:
            parts.append(f"# {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                parts.append(" | ".join("" if cell is None else str(cell) for cell in row))
        return "\n".join(parts)
    finally:
        workbook.close()


def extract_pptx(data: bytes) -> str:
    presentation = Presentation(BytesIO(data))
    parts: list[str] = []
    for index, slide in enumerate(presentation.slides, 1):
        parts.append(f"# Slide {index}")
        for shape in slide.shapes:
            if isinstance(shape, Shape) and shape.has_text_frame:
                parts.append(shape.text)
            elif isinstance(shape, GraphicFrame) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                parts.append(f"Notes: {notes.text}")
    return "\n".join(parts)


def extract_notebook(data: bytes) -> str:
    decoded = cast(object, json.loads(data))
    if not isinstance(decoded, dict):
        raise TypeError("Notebook must be an object")
    notebook = cast(dict[str, object], decoded)
    cells = notebook.get("cells")
    if not isinstance(cells, list):
        raise TypeError("Notebook must contain a cells array")
    parts: list[str] = []
    for item in cast(list[object], cells):
        if not isinstance(item, dict):
            raise TypeError("Notebook cell must be an object")
        cell = cast(dict[str, object], item)
        source = cell.get("source")
        if isinstance(source, str):
            text = source
        elif isinstance(source, list) and all(isinstance(line, str) for line in cast(list[object], source)):
            text = "".join(cast(list[str], source))
        else:
            raise ValueError("Notebook cell source must be text")
        parts.append(f"```\n{text}\n```" if cell.get("cell_type") == "code" else text)
    return "\n\n".join(parts)


def extract_email(data: bytes) -> str:
    message = BytesParser(policy=policy.default).parsebytes(data)
    parts = [f"{name}: {message.get(name, '')}" for name in ("Subject", "From", "To", "Date")]
    for part in message.walk():
        if part.get_content_maintype() == "text" and part.get_content_disposition() != "attachment":
            payload = part.get_payload(decode=True)
            if isinstance(payload, bytes):
                parts.append(payload.decode(part.get_content_charset() or "utf-8", errors="replace"))
    return "\n\n".join(parts)


def extract_archive(data: bytes) -> str:
    """Read supported members without writing archive paths or opening nested archives."""
    parts: list[str] = []
    with ZipFile(BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > MAX_ARCHIVE_ENTRIES or sum(entry.file_size for entry in entries) > MAX_ARCHIVE_BYTES:
            raise ValueError("Archive exceeds document conversion size limits")
        for entry in entries:
            extension = Path(entry.filename).suffix.lower()
            if entry.is_dir() or extension not in SUPPORTED_EXTENSIONS or extension == ".zip":
                continue
            parts.append(f"# {entry.filename}\n\n{extract_content(entry.filename, archive.read(entry))}")
    if not parts:
        raise ValueError("Archive contains no supported documents")
    return "\n\n".join(parts)


def extract_content(name: str, data: bytes) -> str:
    extension = Path(name).suffix.lower()
    if extension in PLAIN_TEXT_EXTENSIONS:
        return data.decode("utf-8", errors="replace")
    if extension == ".docx":
        return extract_docx(data)
    if extension == ".xlsx":
        return extract_xlsx(data)
    if extension == ".pptx":
        return extract_pptx(data)
    if extension == ".pdf":
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted:
            raise ValueError("PDF is encrypted; provide an unlocked copy")
        text = "\n\n".join(page.extract_text() for page in reader.pages)
        if not text.strip():
            raise ValueError("PDF contains no extractable text; OCR is required")
        return text
    if extension == ".ipynb":
        return extract_notebook(data)
    if extension == ".eml":
        return extract_email(data)
    if extension == ".zip":
        return extract_archive(data)
    if extension in IMAGE_EXTENSIONS:
        with Image.open(BytesIO(data)) as image:
            return f"Image: {name}\nFormat: {image.format}\nSize: {image.width} x {image.height}"
    raise ValueError(f"Unsupported document format: {extension or '(none)'}")


def convert_document(path: Path) -> str:
    return extract_content(path.name, path.read_bytes())
