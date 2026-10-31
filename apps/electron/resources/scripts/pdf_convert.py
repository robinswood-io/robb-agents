from pypdf.errors import PyPdfError

"""Native PDF commands, registered on the shared Click group."""

from collections.abc import Sequence
from importlib import import_module
from io import BytesIO
from typing import Protocol, cast

from pdf_common import (
    Path,
    _render_page_to_pil,
    cli,
    click,
    io,
    parse_page_range,
    pdfium,
    sys,
)

_PDF_OPERATION_ERRORS = (OSError, ValueError, TypeError, RuntimeError, KeyError, IndexError, ImportError, PyPdfError)

class _DocxDocumentLike(Protocol):
    def add_paragraph(self, text: str) -> object: ...

    def add_page_break(self) -> object: ...

    def save(self, path: str) -> None: ...


class _PictureCollectionLike(Protocol):
    def add_picture(self, image: BytesIO, left: object, top: object, width: object, height: object) -> object: ...


class _SlideLike(Protocol):
    shapes: _PictureCollectionLike


class _SlideCollectionLike(Protocol):
    def add_slide(self, layout: object) -> _SlideLike: ...


class _PresentationLike(Protocol):
    slide_width: int
    slide_height: int
    slide_layouts: Sequence[object]
    slides: _SlideCollectionLike

    def save(self, path: str) -> None: ...


def _as_presentation(value: object) -> _PresentationLike:
    """Adapt python-pptx after checking its runtime presentation surface."""
    if not callable(getattr(value, "save", None)):
        raise TypeError("Presentation does not expose save")
    if not all(hasattr(value, name) for name in ("slide_width", "slide_height", "slide_layouts", "slides")):
        raise TypeError("Presentation is missing slide collections")
    return cast(_PresentationLike, value)


def _as_docx_document(value: object) -> _DocxDocumentLike:
    """Adapt python-docx after checking its runtime document surface."""
    if not all(callable(getattr(value, name, None)) for name in ("add_paragraph", "add_page_break", "save")):
        raise TypeError("Document does not implement the python-docx conversion interface")
    return cast(_DocxDocumentLike, value)



@cli.command("to-image")
@click.argument("file", type=click.Path(exists=True, dir_okay=False))
@click.option("--format", "fmt", type=click.Choice(["png", "jpg"]), default="png", help="Image format (default: png).")
@click.option("--dpi", type=int, default=150, help="Resolution in DPI (default: 150).")
@click.option("--pages", type=str, default=None, help="Page range to convert (default: all).")
@click.option("-o", "--output", type=click.Path(), required=True, help="Output directory.")
def to_image(file: str, fmt: str, dpi: int, pages: str | None, output: str) -> None:
    """Convert PDF pages to images (PNG or JPG)."""
    try:
        pdf = pdfium.PdfDocument(file)
        total = len(pdf)
        indices = parse_page_range(pages, total) if pages else list(range(total))
        out_dir = Path(output)
        out_dir.mkdir(parents=True, exist_ok=True)

        for idx in indices:
            page = pdf[idx]
            img = _render_page_to_pil(page, dpi=dpi)
            ext = "jpg" if fmt == "jpg" else "png"
            out_path = out_dir / f"page_{idx + 1}.{ext}"
            if fmt == "jpg":
                img.save(str(out_path), "JPEG", quality=85)
            else:
                img.save(str(out_path), "PNG")
            click.echo(f"Saved {out_path}", err=True)

        click.echo(f"Converted {len(indices)} pages to {fmt.upper()} in {output}", err=True)
    except _PDF_OPERATION_ERRORS as e:
        click.echo(f"Error: {e}", err=True)
        sys.exit(1)


@cli.command("from-image")
@click.argument("files", nargs=-1, required=True, type=click.Path(exists=True, dir_okay=False))
@click.option("-o", "--output", type=click.Path(), required=True, help="Output PDF file path.")
def from_image(files: tuple[str, ...], output: str) -> None:
    """Convert images to a PDF document.

    Each image becomes one page. Supports PNG, JPG, BMP, GIF, TIFF, WebP.
    """
    try:
        module = import_module("img2pdf")
        converter: object = getattr(module, "convert", None)
        if not callable(converter):
            raise TypeError("img2pdf does not expose convert")
        img_data: list[bytes] = [Path(img_path).read_bytes() for img_path in files]
        pdf_bytes: object = converter(img_data)
        if not isinstance(pdf_bytes, bytes):
            raise TypeError("img2pdf.convert did not return bytes")
        _ = Path(output).write_bytes(pdf_bytes)
        click.echo(f"PDF written to {output} ({len(files)} pages from images)", err=True)
    except _PDF_OPERATION_ERRORS as e:
        click.echo(f"Error: {e}", err=True)
        sys.exit(1)


@cli.command("to-docx")
@click.argument("file", type=click.Path(exists=True, dir_okay=False))
@click.option("--pages", type=str, default=None, help="Page range to convert (default: all).")
@click.option("-o", "--output", type=click.Path(), required=True, help="Output DOCX file path.")
def to_docx(file: str, pages: str | None, output: str) -> None:
    """Convert PDF to Word document (text extraction based)."""
    from docx import Document as DocxDocument

    try:
        pdf = pdfium.PdfDocument(file)
        total = len(pdf)
        indices = parse_page_range(pages, total) if pages else list(range(total))

        docx_doc = _as_docx_document(cast(object, DocxDocument()))
        for i, idx in enumerate(indices):
            page = pdf[idx]
            textpage = page.get_textpage()
            text = textpage.get_text_bounded()
            for line in text.splitlines():
                if line.strip():
                    _ = docx_doc.add_paragraph(line.strip())
            if i < len(indices) - 1:
                _ = docx_doc.add_page_break()

        docx_doc.save(output)
        click.echo(f"DOCX written to {output}", err=True)
    except _PDF_OPERATION_ERRORS as e:
        click.echo(f"Error: {e}", err=True)
        sys.exit(1)


@cli.command("to-pptx")
@click.argument("file", type=click.Path(exists=True, dir_okay=False))
@click.option("--dpi", type=int, default=200, help="Image resolution for slides (default: 200).")
@click.option("--pages", type=str, default=None, help="Page range to convert (default: all).")
@click.option("-o", "--output", type=click.Path(), required=True, help="Output PPTX file path.")
def to_pptx(file: str, dpi: int, pages: str | None, output: str) -> None:
    """Convert PDF pages to PowerPoint slides (image-based)."""
    from pptx import Presentation
    from pptx.util import Emu, Inches

    try:
        pdf = pdfium.PdfDocument(file)
        total = len(pdf)
        indices = parse_page_range(pages, total) if pages else list(range(total))

        if not indices:
            click.echo("Error: no pages selected for conversion.", err=True)
            sys.exit(1)

        prs = _as_presentation(cast(object, Presentation()))

        # Set slide size to match first page aspect ratio (points -> EMU: 1pt = 12700 EMU)
        first_page = pdf[indices[0]]
        first_w, first_h = first_page.get_size()
        prs.slide_width = Emu(int(first_w * 12700))
        prs.slide_height = Emu(int(first_h * 12700))
        blank_layout = prs.slide_layouts[6]  # blank slide

        for idx in indices:
            page = pdf[idx]
            img = _render_page_to_pil(page, dpi=dpi)

            img_buffer = io.BytesIO()
            img.save(img_buffer, format="PNG")
            _ = img_buffer.seek(0)

            slide = prs.slides.add_slide(blank_layout)
            _ = slide.shapes.add_picture(
                img_buffer, Inches(0), Inches(0),
                prs.slide_width, prs.slide_height,
            )

        prs.save(output)
        click.echo(f"PPTX written to {output} ({len(indices)} slides)", err=True)
    except _PDF_OPERATION_ERRORS as e:
        click.echo(f"Error: {e}", err=True)
        sys.exit(1)
